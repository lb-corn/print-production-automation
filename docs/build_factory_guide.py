from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

root=Path(__file__).resolve().parents[1]
doc=Document(); section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=Inches(.7);section.bottom_margin=Inches(.7);section.left_margin=Inches(.8);section.right_margin=Inches(.8)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',25),('Heading 1',15),('Heading 2',12)]:
    style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('000000')
    style.paragraph_format.space_before=Pt(14);style.paragraph_format.space_after=Pt(7)
    style.paragraph_format.keep_with_next=True
header=section.header.paragraphs[0];header.text='LUCKY BAGS MANUFACTURING';header.style='Caption';header.runs[0].font.color.rgb=RGBColor.from_string('637987')
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Factory Agent guide  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for run in footer.runs:run.font.size=Pt(9)

def text_inline(paragraph,text):
    # Plain, editable body text; remove Markdown delimiters from prose.
    parts=re.split(r'(\*\*.*?\*\*|`[^`]+`)',text)
    for part in parts:
        run=paragraph.add_run(part.strip('*`'))
        if part.startswith('**'):run.bold=True
        if part.startswith('`'):run.font.name='Consolas';run.font.size=Pt(9)

lines=(root/'automation/README.md').read_text(encoding='utf-8-sig').splitlines()
i=0;code=False
while i<len(lines):
    line=lines[i];i+=1
    if line.startswith('```'):code=not code;continue
    if not line.strip():continue
    if code:
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2)
        r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(9)
        continue
    if line.startswith('# '):
        doc.add_paragraph('Lucky Bags Factory Agent Guide',style='Title');continue
    if line.startswith('## '):
        if line[3:] in ['Deploy the web application first','Required configuration','Controlled pilot','Continuous operation','Updates and rollback']:
            doc.add_page_break()
        doc.add_paragraph(re.sub(r'[^A-Za-z0-9 ]','',line[3:]),style='Heading 1');continue
    if line.startswith('|'):
        tablelines=[line]
        while i<len(lines) and lines[i].startswith('|'):tablelines.append(lines[i]);i+=1
        parsed=[[x.strip() for x in row.strip('|').split('|')] for row in tablelines if not re.fullmatch(r'[| :\-]+',row)]
        table=doc.add_table(rows=1,cols=len(parsed[0]));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        for j,width in enumerate([2.1,4.8]):table.columns[j].width=Inches(width)
        for rindex,row in enumerate(parsed):
            cells=table.rows[0].cells if rindex==0 else table.add_row().cells
            for j,value in enumerate(row):
                cell=cells[j];cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                text_inline(cell.paragraphs[0],value)
                for run in cell.paragraphs[0].runs:run.font.size=Pt(10);run.bold=rindex==0
                cell.paragraphs[0].paragraph_format.space_after=Pt(5);cell.paragraphs[0].paragraph_format.space_before=Pt(5)
                tcPr=cell._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
                tcPr.append(borders)
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EFF3' if rindex==0 else ('F7F9FA' if rindex%2==0 else 'FFFFFF'));tcPr.append(shade)
            if rindex==0:
                repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
        doc.add_paragraph();continue
    match=re.match(r'^(\d+)\. (.*)',line)
    if line.startswith('- '):p=doc.add_paragraph(style='List Bullet');text_inline(p,line[2:])
    elif match:p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.first_line_indent=Inches(-.2);text_inline(p,match[1]+'. '+match[2])
    else:p=doc.add_paragraph();text_inline(p,line)

for border in list(doc.styles.element.xpath('.//w:pBdr')) + list(doc.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)

doc.core_properties.title='Lucky Bags Factory Agent Guide'
doc.core_properties.subject='Secure installation and operator controlled Epson Edge Print workflow'
doc.core_properties.author='Lucky Bags Manufacturing'
doc.save(root/'docs/LuckyBags-Factory-Guide.docx')
print('Created Word guide')
